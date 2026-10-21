from __future__ import annotations

import csv
import io
import logging
from typing import Any

from openpyxl import Workbook

logger = logging.getLogger("dhis2_analyst.xlsx_gen")


def data_to_xlsx(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    headers = payload.get("headers") or []
    rows = payload.get("rows") or []
    logger.info("xlsx_generate_start", extra={"row_count": len(rows), "header_count": len(headers)})
    if headers:
        ws.append(headers)
    for row in rows:
        # DHIS2 analytics returns values as strings; write numbers as numbers
        # so the sheet can be summed/charted in Excel.
        ws.append([_as_number(v) for v in row])

    meta = wb.create_sheet("Metadata")
    for key, value in _readable_metadata(payload.get("metadata") or {}):
        meta.append([key, value])

    sources = wb.create_sheet("Sources")
    sources.append(["title", "url", "confidence"])
    for item in payload.get("sources") or []:
        sources.append([item.get("title", ""), item.get("url", ""), item.get("confidence", "")])

    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    logger.info("xlsx_generate_ok", extra={"size_bytes": len(data), "row_count": len(rows), "sheet_count": len(wb.worksheets)})
    return data


def _as_number(value: Any) -> Any:
    if isinstance(value, str):
        try:
            num = float(value)
        except ValueError:
            return value
        return int(num) if num.is_integer() and "." not in value else num
    return value


def _readable_metadata(metadata: dict[str, Any]) -> list[tuple[str, str]]:
    """Human-readable provenance (names + UIDs) instead of a raw dict dump."""
    out: list[tuple[str, str]] = []
    raw = metadata.get("raw_metadata") or {}
    items = raw.get("items") or {}
    dims = raw.get("dimensions") or {}

    def label(uid: str) -> str:
        name = (items.get(uid) or {}).get("name")
        return f"{name} ({uid})" if name and name != uid else str(uid)

    if metadata.get("data_source"):
        out.append(("Source", f"DHIS2 {metadata['data_source']}"))
    for key, title in (("dx", "Data (indicator / data element UID)"), ("ou", "Organisation units"), ("pe", "Periods")):
        if dims.get(key):
            out.append((title, "; ".join(label(u) for u in dims[key])))
    params = metadata.get("analytics_params") or {}
    if params.get("dimension"):
        out.append(("Analytics query", "&".join(f"dimension={d}" for d in params["dimension"])))
    for key, value in metadata.items():
        if key not in {"raw_metadata", "data_source", "analytics_params"}:
            out.append((str(key), str(value)))
    return out


def data_to_csv(payload: dict[str, Any]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    headers = payload.get("headers") or []
    rows = payload.get("rows") or []
    logger.info("csv_generate_start", extra={"row_count": len(rows), "header_count": len(headers)})
    if headers:
        writer.writerow(headers)
    writer.writerows(rows)
    data = out.getvalue().encode("utf-8")
    logger.info("csv_generate_ok", extra={"size_bytes": len(data), "row_count": len(rows)})
    return data
