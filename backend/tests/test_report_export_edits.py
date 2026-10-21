"""Report exports use the user's edited report HTML, safely.

The browser now posts the report editor's current HTML (or the original
generated HTML when nothing was edited) to /api/generate/{pdf,docx}. These tests
pin down:
- edited text (new title, added sentence) reaches the DOCX and the PDF;
- formatting survives: headings, lists, tables, the data-sources section;
- unedited, clean report HTML goes to the generators byte-for-byte unchanged;
- script/style/iframe/img and event-handler / javascript: injection is removed,
  and WeasyPrint never fetches external resources.
"""
import asyncio
import io
import re
from pathlib import Path

import pytest
from docx import Document

from backend.config import Settings
from backend.app.agent.renderers.report import render_report
from backend.app.generators import pdf_gen
from backend.app.generators.docx_gen import html_to_docx
from backend.app.generators.html_sanitize import report_text, sanitize_report_html
from backend.app.models import Identity, ReportPayload


def _state():
    return {
        "dhis2_data": {
            "headers": ["Organisation Unit", "Period", "Value"],
            "rows": [
                ["Sierra Leone", "202501", "80.5"],
                ["Sierra Leone", "202502", "95.25"],
                ["Sierra Leone", "202503", "101.0"],
            ],
            "metadata": {"data_source": "analytics_api"},
        },
        "metrics": [{"label": "ANC 1 Coverage", "uid": "Uvn6LCg7dVU"}],
        "org_unit": {"label": "Sierra Leone"},
        "periods": ["202501", "202502", "202503"],
        "web_context": [{"url": "https://example.org/anc", "title": "ANC guidance & notes"}],
    }


GENERATED = render_report(_state())

# What the Tiptap editor produces after a user changes the title and adds a
# sentence (Tiptap re-serialises: <tbody> only, <p> inside cells).
EDITED = (
    "<h1>Edited title for the ANC review</h1>"
    "<p><em>Generated 21 October 2026 | Periods: 202501, 202502, 202503</em></p>"
    "<h2>Executive Summary</h2>"
    "<p>This report analyses ANC 1 Coverage for Sierra Leone. "
    "This sentence was added by the analyst.</p>"
    "<h2>Key Findings</h2>"
    "<ul><li><p>Average: <strong>92.3</strong></p></li><li><p>Range: 80 to 101</p></li></ul>"
    "<h2>Data Table</h2>"
    "<table><tbody>"
    "<tr><th><p>Organisation Unit</p></th><th><p>Period</p></th><th><p>Value</p></th></tr>"
    "<tr><td><p>Sierra Leone</p></td><td><p>202501</p></td><td><p>80.5</p></td></tr>"
    "</tbody></table>"
    "<h2>Recommendations</h2>"
    "<ol><li><p>Discuss findings with programme teams.</p></li></ol>"
    "<h2>Data Sources</h2>"
    '<ul><li><p>DHIS2 Analytics (analytics_api) — 1 indicator(s) queried</p></li></ul>'
)

MALICIOUS = (
    "<h1>Title</h1>"
    "<script>alert('x')</script>"
    "<style>body{background:url(http://169.254.169.254/latest)}</style>"
    '<p onclick="steal()" style="color:red">Visible text</p>'
    '<img src="file:///etc/passwd" onerror="x()">'
    '<iframe src="http://internal.local/">iframe body</iframe>'
    '<p><a href="javascript:alert(1)">bad link</a> <a href="https://dhis2.org">good link</a></p>'
    '<p style="background:url(file:///etc/shadow)">styled</p>'
    "<!-- hidden comment -->"
    '<svg><script>alert(2)</script></svg>'
    "<blink>unknown tag text</blink>"
)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _docx_styles(data: bytes) -> list[str]:
    return [p.style.name for p in Document(io.BytesIO(data)).paragraphs]


# --- sanitiser ---------------------------------------------------------------

def test_clean_generated_report_passes_through_unchanged():
    # Unedited downloads must be exactly what they were before this change.
    assert sanitize_report_html(GENERATED) == GENERATED


def test_clean_editor_html_passes_through_unchanged():
    assert sanitize_report_html(EDITED) == EDITED


@pytest.mark.parametrize("html", [
    '<p class="cached-banner"><strong>Cached</strong> — snapshot</p><h1>T</h1>',
    '<table border="1"><thead><tr><th scope="col" colspan="2">A</th></tr></thead></table>',
    '<p style="text-align:center">centred</p>',
    '<a href="#sec">anchor</a> <a href="mailto:a@b.org">mail</a>',
])
def test_other_clean_markup_unchanged(html):
    assert sanitize_report_html(html) == html


def test_malicious_markup_is_removed():
    out = sanitize_report_html(MALICIOUS)
    low = out.lower()
    for bad in ("<script", "alert(", "<style", "<iframe", "iframe body", "<img", "passwd",
                "onclick", "onerror", "javascript:", "url(", "<svg", "<!--", "hidden comment",
                "169.254", "<blink"):
        assert bad not in low, bad
    # Legitimate content and formatting stay.
    assert "<h1>Title</h1>" in out
    assert "Visible text" in out and 'style="color:red"' in out
    assert '<a href="https://dhis2.org">good link</a>' in out
    assert "bad link" in out          # text kept, dangerous href dropped
    assert "unknown tag text" in out  # unknown tag unwrapped, text kept
    assert "styled" in out


def test_sanitize_empty_and_non_string():
    assert sanitize_report_html("") == ""
    assert sanitize_report_html(None) == ""  # type: ignore[arg-type]


# --- DOCX ----------------------------------------------------------------------

def test_docx_contains_the_users_edits_and_keeps_formatting():
    data = html_to_docx(EDITED, "DHIS2 Report")
    text = _docx_text(data)
    assert "Edited title for the ANC review" in text
    assert "This sentence was added by the analyst." in text
    assert "Data Sources" in text and "DHIS2 Analytics (analytics_api)" in text
    assert "Sierra Leone" in text and "80.5" in text  # table cells
    styles = _docx_styles(data)
    assert "Heading 1" in styles and "Heading 2" in styles
    assert "List Bullet" in styles and "List Number" in styles
    assert len(Document(io.BytesIO(data)).tables) == 1


def test_docx_unedited_report_text_identical_to_before():
    # Same input string as before -> same document text (sanitiser is a no-op).
    data = html_to_docx(GENERATED, "DHIS2 Report")
    text = _docx_text(data)
    assert "ANC 1 Coverage Review — Sierra Leone" in text
    assert "ANC guidance &amp; notes" not in text  # entities decoded as before
    assert "ANC guidance & notes" in text


def test_docx_drops_injected_script_text():
    text = _docx_text(html_to_docx(MALICIOUS, "DHIS2 Report"))
    assert "alert" not in text
    assert "Visible text" in text


# --- PDF -----------------------------------------------------------------------

class _CaptureHTML:
    calls: list[dict] = []

    def __init__(self, string=None, url_fetcher=None, **kw):
        _CaptureHTML.calls.append({"string": string, "url_fetcher": url_fetcher})

    def write_pdf(self):
        return b"%PDF-1.7 captured"


def _capture(monkeypatch):
    import weasyprint
    _CaptureHTML.calls = []
    monkeypatch.setattr(weasyprint, "HTML", _CaptureHTML)
    return _CaptureHTML.calls


def test_pdf_receives_the_users_edits(monkeypatch):
    calls = _capture(monkeypatch)
    pdf_gen.html_to_pdf(EDITED)
    doc = calls[0]["string"]
    assert "Edited title for the ANC review" in doc
    assert "This sentence was added by the analyst." in doc
    assert "<h2>Data Sources</h2>" in doc and "<table>" in doc and "<ol>" in doc


def test_pdf_unedited_report_input_is_unchanged(monkeypatch):
    calls = _capture(monkeypatch)
    pdf_gen.html_to_pdf(GENERATED)
    assert calls[0]["string"] == pdf_gen._wrap(GENERATED)


def test_pdf_injection_removed_and_fetching_blocked(monkeypatch):
    calls = _capture(monkeypatch)
    pdf_gen.html_to_pdf(MALICIOUS)
    body = calls[0]["string"].split("<body>", 1)[1].lower()
    assert "<script" not in body and "<img" not in body and "url(" not in body
    fetcher = calls[0]["url_fetcher"]
    assert fetcher is not None
    with pytest.raises(Exception):
        fetcher("file:///etc/passwd")


def test_pdf_fallback_also_sanitised(monkeypatch):
    import weasyprint

    def boom(*a, **k):
        raise RuntimeError("no weasyprint")

    monkeypatch.setattr(weasyprint, "HTML", boom)
    out = pdf_gen.html_to_pdf(MALICIOUS)
    assert out.startswith(b"%PDF-")
    assert b"alert" not in out and b"Visible text" in out


def test_real_pdf_text_contains_edits():
    pypdf = pytest.importorskip("pypdf")
    data = pdf_gen.html_to_pdf(EDITED)
    text = "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(io.BytesIO(data)).pages)
    text = re.sub(r"\s+", " ", text)
    assert "Edited title for the ANC review" in text
    assert "This sentence was added by the analyst." in text
    assert "Data Sources" in text


# --- endpoints -----------------------------------------------------------------

def _run(coro):
    return asyncio.new_event_loop().run_until_complete(coro)


def test_generate_endpoints_use_posted_edited_html(tmp_path):
    from backend.main import generate_docx, generate_pdf

    settings = Settings(jwt_secret="x" * 32, temp_file_dir=str(tmp_path))
    identity = Identity(user_id="tester")
    payload = ReportPayload(html=EDITED + MALICIOUS, title="DHIS2 Report")

    docx_meta = _run(generate_docx(payload, settings, identity))
    docx_bytes = (Path(settings.temp_path) / docx_meta["file_id"]).read_bytes()
    text = _docx_text(docx_bytes)
    assert "Edited title for the ANC review" in text and "alert" not in text

    pdf_meta = _run(generate_pdf(payload, settings, identity))
    pdf_bytes = (Path(settings.temp_path) / pdf_meta["file_id"]).read_bytes()
    assert pdf_bytes.startswith(b"%PDF-")


def test_report_text_helper():
    assert report_text("<h1>A</h1><p>b <em>c</em></p>") == "A b c"
